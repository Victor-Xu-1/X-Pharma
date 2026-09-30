from __future__ import annotations

import argparse
import hashlib
import json
import os
import tempfile
from datetime import UTC, datetime
from importlib import import_module
from pathlib import Path
from typing import Any

import gemmi
import httpx
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from pharma_intel.config import get_settings
from pharma_intel.ingest.parser_client import build_document_parser

Chem: Any = import_module("rdkit.Chem")


def _write_fixtures(root: Path) -> list[Path]:
    markdown = root / "evidence.md"
    markdown.write_text("EGFR L858R evidence", encoding="utf-8")

    html = root / "evidence.html"
    html.write_text("<html><body><h1>KRAS evidence</h1></body></html>", encoding="utf-8")

    document_path = root / "evidence.docx"
    document = Document()
    document.add_paragraph("HER2 clinical evidence")
    document.save(str(document_path))

    presentation_path = root / "evidence.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "BTK competitive landscape"
    presentation.save(str(presentation_path))

    workbook_path = root / "evidence.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet["A1"] = "compound"
    worksheet["B1"] = "IC50"
    worksheet["A2"] = "CMP-001"
    worksheet["B2"] = 12.5
    workbook.save(workbook_path)
    workbook.close()

    pdf_path = root / "evidence.pdf"
    pdf = PdfWriter()
    page = pdf.add_blank_page(width=612, height=792)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): pdf._add_object(font)})}
    )
    content = DecodedStreamObject()
    content.set_data(b"BT /F1 12 Tf 72 720 Td (EGFR PDF evidence) Tj ET")
    page[NameObject("/Contents")] = pdf._add_object(content)
    with pdf_path.open("wb") as output:
        pdf.write(output)

    molecule = Chem.MolFromSmiles("CCO")
    if molecule is None:
        raise RuntimeError("RDKit probe fixture creation failed")
    molecule.SetProp("_Name", "Ethanol")
    sdf_path = root / "ethanol.sdf"
    sdf_writer = Chem.SDWriter(str(sdf_path))
    sdf_writer.write(molecule)
    sdf_writer.close()
    mol_path = root / "ethanol.mol"
    Chem.MolToMolFile(molecule, str(mol_path))

    pdb_path = root / "complex.pdb"
    pdb_path.write_text(
        "HEADER    PARSER PROBE\n"
        "ATOM      1  N   GLY A   1      11.104  13.207  10.000  1.00 20.00           N\n"
        "ATOM      2  CA  GLY A   1      12.000  12.100  10.000  1.00 20.00           C\n"
        "HETATM    3  C1  LIG A 101      14.000  11.000  10.000  1.00 20.00           C\n"
        "END\n",
        encoding="ascii",
    )
    structure = gemmi.read_structure(str(pdb_path))
    cif_path = root / "complex.cif"
    structure.make_mmcif_document().write_file(str(cif_path))
    mmcif_path = root / "complex.mmcif"
    structure.make_mmcif_document().write_file(str(mmcif_path))

    return [
        markdown,
        html,
        document_path,
        presentation_path,
        workbook_path,
        pdf_path,
        sdf_path,
        mol_path,
        pdb_path,
        cif_path,
        mmcif_path,
    ]


def _protocol_rejections(base_url: str, token: str) -> tuple[int, int]:
    payload = b"protocol boundary evidence"
    digest = hashlib.sha256(payload).hexdigest()
    headers = {
        "Content-Length": str(len(payload)),
        "Content-Type": "application/octet-stream",
        "X-Content-SHA256": digest,
    }
    with httpx.Client(timeout=10, trust_env=False) as client:
        unauthorized = client.post(
            f"{base_url.rstrip('/')}/internal/v1/parse",
            params={"filename": "unauthorized.md", "max_chars": 100_000},
            headers={**headers, "Authorization": "Bearer invalid-parser-token"},
            content=payload,
        )
        digest_mismatch = client.post(
            f"{base_url.rstrip('/')}/internal/v1/parse",
            params={"filename": "digest.md", "max_chars": 100_000},
            headers={**headers, "Authorization": f"Bearer {token}", "X-Content-SHA256": "0" * 64},
            content=payload,
        )
    return unauthorized.status_code, digest_mismatch.status_code


def probe() -> dict[str, object]:
    settings = get_settings()
    if settings.parser_backend != "service":
        raise RuntimeError("Parser acceptance requires PARSER_BACKEND=service")
    parser = build_document_parser(settings)
    documents: list[dict[str, object]] = []
    with tempfile.TemporaryDirectory(prefix="pharma-parser-probe-") as temporary_directory:
        for path in _write_fixtures(Path(temporary_directory)):
            parsed = parser.parse(path, settings.parser_max_text_chars)
            documents.append(
                {
                    "suffix": path.suffix,
                    "parser_name": parsed.parser_name,
                    "parser_version": parsed.parser_version,
                    "text_sha256": hashlib.sha256(parsed.text.encode("utf-8")).hexdigest(),
                    "metadata_keys": sorted(parsed.metadata),
                }
            )
    unauthorized_status, digest_mismatch_status = _protocol_rejections(
        settings.parser_service_url,
        settings.parser_service_token,
    )
    if unauthorized_status != 401 or digest_mismatch_status != 400:
        raise RuntimeError("Parser service protocol rejection controls did not pass")
    return {
        "schema_version": 1,
        "status": "passed",
        "generated_at": datetime.now(UTC).isoformat(),
        "production_claim": False,
        "parser_backend": settings.parser_backend,
        "document_count": len(documents),
        "documents": documents,
        "unauthorized_status": unauthorized_status,
        "digest_mismatch_status": digest_mismatch_status,
    }


def _write_atomic(path: Path, payload: bytes) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(f".{path.name}.{os.getpid()}.tmp")
    descriptor = os.open(temporary, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
    try:
        with os.fdopen(descriptor, "wb") as output:
            output.write(payload)
            output.flush()
            os.fsync(output.fileno())
        try:
            os.link(temporary, path, follow_symlinks=False)
        except FileExistsError as exc:
            raise RuntimeError(f"Refusing to overwrite parser probe evidence: {path}") from exc
    finally:
        temporary.unlink(missing_ok=True)


def run() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    payload = (json.dumps(probe(), indent=2, sort_keys=True) + "\n").encode()
    if args.output:
        _write_atomic(args.output, payload)
    print(payload.decode(), end="")
