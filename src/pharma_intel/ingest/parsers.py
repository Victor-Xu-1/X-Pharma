from __future__ import annotations

import csv
import importlib.metadata
import json
import stat
import zipfile
from dataclasses import dataclass
from pathlib import Path, PurePosixPath
from typing import Any, Protocol

from bs4 import BeautifulSoup
from docx import Document
from lxml import etree  # type: ignore[import-untyped]  # lxml does not publish PEP 561 metadata.
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader


class DocumentParseError(RuntimeError):
    pass


class UnsupportedDocumentError(DocumentParseError):
    pass


class DocumentOcrRequired(DocumentParseError):
    """The document is valid but has no machine-readable text layer."""


SCIENTIFIC_SUFFIXES = frozenset({".cif", ".mmcif", ".mol", ".pdb", ".sdf"})
OFFICE_MAX_ARCHIVE_MEMBERS = 4096
OFFICE_MAX_UNCOMPRESSED_BYTES = 536_870_912
OFFICE_MAX_MEMBER_UNCOMPRESSED_BYTES = 268_435_456
OFFICE_MAX_COMPRESSION_RATIO = 200
OFFICE_COMPRESSION_RATIO_MIN_BYTES = 1_048_576


@dataclass(frozen=True)
class ParsedDocument:
    text: str
    metadata: dict[str, Any]
    parser_name: str
    parser_version: str


class Parser(Protocol):
    def __call__(self, path: Path, max_chars: int, /) -> ParsedDocument: ...


class TextAccumulator:
    def __init__(self, max_chars: int) -> None:
        self.max_chars = max_chars
        self.parts: list[str] = []
        self.length = 0

    def add(self, value: str) -> None:
        if not value:
            return
        if self.length + len(value) > self.max_chars:
            raise DocumentParseError(f"Extracted text exceeds the configured {self.max_chars} character limit")
        self.parts.append(value)
        self.length += len(value)

    def render(self) -> str:
        return "\n".join(self.parts).strip()


def parse_document(path: Path, max_chars: int) -> ParsedDocument:
    suffix = path.suffix.casefold()
    parser = PARSERS.get(suffix)
    if parser is None and suffix in SCIENTIFIC_SUFFIXES:
        _register_scientific_parsers()
        parser = PARSERS.get(suffix)
    if parser is None:
        raise UnsupportedDocumentError(f"No parser is registered for {suffix or 'extensionless files'}")
    try:
        parsed = parser(path, max_chars)
    except DocumentParseError:
        raise
    except Exception as exc:
        raise DocumentParseError(f"{path.name} could not be parsed: {exc}") from exc
    if not parsed.text.strip():
        raise DocumentParseError(f"{path.name} produced no extractable text")
    return parsed


def _parse_pdf(path: Path, max_chars: int) -> ParsedDocument:
    reader = PdfReader(path, strict=True)
    if reader.is_encrypted:
        raise DocumentParseError("Encrypted PDF files require a configured decryption workflow")
    page_texts = [(page.extract_text() or "") for page in reader.pages]
    if not any(text.strip() for text in page_texts):
        raise DocumentOcrRequired("PDF has no machine-readable text layer and requires OCR")
    output = TextAccumulator(max_chars)
    for index, text in enumerate(page_texts, start=1):
        output.add(f"[[page:{index}]]")
        output.add(text)
    return ParsedDocument(
        output.render(),
        {
            "page_count": len(page_texts),
            "machine_readable_page_count": sum(bool(text.strip()) for text in page_texts),
        },
        "pypdf",
        importlib.metadata.version("pypdf"),
    )


def _parse_ocr_image(_path: Path, _max_chars: int) -> ParsedDocument:
    raise DocumentOcrRequired("Document image requires OCR")


def _parse_docx(path: Path, max_chars: int) -> ParsedDocument:
    _validate_office_archive(path)
    document = Document(str(path))
    output = TextAccumulator(max_chars)
    for paragraph in document.paragraphs:
        output.add(paragraph.text)
    for table_index, table in enumerate(document.tables, start=1):
        output.add(f"[[table:{table_index}]]")
        for row in table.rows:
            output.add("\t".join(cell.text for cell in row.cells))
    return ParsedDocument(
        output.render(),
        {"paragraph_count": len(document.paragraphs), "table_count": len(document.tables)},
        "python-docx",
        importlib.metadata.version("python-docx"),
    )


def _parse_pptx(path: Path, max_chars: int) -> ParsedDocument:
    _validate_office_archive(path)
    presentation = Presentation(str(path))
    output = TextAccumulator(max_chars)
    for slide_index, slide in enumerate(presentation.slides, start=1):
        output.add(f"[[slide:{slide_index}]]")
        for shape in slide.shapes:
            text = getattr(shape, "text", "")
            if isinstance(text, str):
                output.add(text)
        if slide.has_notes_slide:
            notes_frame = slide.notes_slide.notes_text_frame
            if notes_frame is not None:
                output.add("[[notes]]")
                output.add(notes_frame.text)
    return ParsedDocument(
        output.render(),
        {"slide_count": len(presentation.slides)},
        "python-pptx",
        importlib.metadata.version("python-pptx"),
    )


def _parse_xlsx(path: Path, max_chars: int) -> ParsedDocument:
    _validate_office_archive(path)
    workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    output = TextAccumulator(max_chars)
    non_empty_cells = 0
    try:
        for worksheet in workbook.worksheets:
            output.add(f"[[sheet:{worksheet.title}]]")
            for row_number, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
                values = ["" if value is None else _cell_text(value) for value in row]
                if not any(values):
                    continue
                non_empty_cells += sum(bool(value) for value in values)
                output.add(f"{row_number}\t" + "\t".join(values))
    finally:
        workbook.close()
    return ParsedDocument(
        output.render(),
        {"sheet_count": len(workbook.sheetnames), "non_empty_cell_count": non_empty_cells},
        "openpyxl",
        importlib.metadata.version("openpyxl"),
    )


def _cell_text(value: Any) -> str:
    if isinstance(value, dict | list | tuple):
        return json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    return str(value).replace("\r", " ").replace("\n", " ")


def _parse_delimited(path: Path, max_chars: int) -> ParsedDocument:
    text, encoding = _read_text(path, max_chars)
    dialect = csv.Sniffer().sniff(text[:16_384], delimiters=",\t;|")
    output = TextAccumulator(max_chars)
    for row_number, row in enumerate(csv.reader(text.splitlines(), dialect), start=1):
        output.add(f"{row_number}\t" + "\t".join(value.strip() for value in row))
    return ParsedDocument(output.render(), {"encoding": encoding}, "python-csv", "stdlib")


def _parse_text(path: Path, max_chars: int) -> ParsedDocument:
    text, encoding = _read_text(path, max_chars)
    return ParsedDocument(text, {"encoding": encoding}, "text", "1")


def _parse_json(path: Path, max_chars: int) -> ParsedDocument:
    text, encoding = _read_text(path, max_chars)
    try:
        value = json.loads(text)
    except json.JSONDecodeError as exc:
        raise DocumentParseError("JSON document is not syntactically valid") from exc
    canonical = json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    if len(canonical) > max_chars:
        raise DocumentParseError(f"Extracted text exceeds the configured {max_chars} character limit")
    return ParsedDocument(
        canonical,
        {"encoding": encoding, "top_level_type": type(value).__name__},
        "python-json",
        "stdlib",
    )


def _parse_html(path: Path, max_chars: int) -> ParsedDocument:
    text, encoding = _read_text(path, max_chars)
    # The stdlib HTMLParser is intentionally avoided for untrusted source files.
    soup = BeautifulSoup(text, "lxml")
    for element in soup(["script", "style", "noscript", "template"]):
        element.decompose()
    output = "\n".join(line.strip() for line in soup.get_text("\n").splitlines() if line.strip())
    if len(output) > max_chars:
        raise DocumentParseError(f"Extracted text exceeds the configured {max_chars} character limit")
    return ParsedDocument(
        output,
        {"encoding": encoding, "title": soup.title.string.strip() if soup.title and soup.title.string else None},
        "beautifulsoup4+lxml",
        importlib.metadata.version("beautifulsoup4"),
    )


def _normalized_xml_text(element: Any) -> str:
    return " ".join(" ".join(element.itertext()).split())


def _first_xpath_text(element: Any, expression: str) -> str | None:
    matches = element.xpath(expression)
    if not matches:
        return None
    text = _normalized_xml_text(matches[0])
    return text or None


def _add_jats_section(output: TextAccumulator, section: Any) -> None:
    title = _first_xpath_text(section, "./*[local-name()='title']")
    section_id = section.get("id")
    marker = title or section_id or "untitled"
    output.add(f"[[section:{marker}]]")
    for child in section:
        local_name = etree.QName(child).localname
        if local_name == "sec":
            _add_jats_section(output, child)
        elif local_name in {"p", "list", "table-wrap", "disp-quote", "boxed-text"}:
            output.add(_normalized_xml_text(child))


def _parse_jats_xml(path: Path, max_chars: int) -> ParsedDocument:
    data = path.read_bytes()
    if len(data) > max(max_chars * 4, 16_777_216):
        raise DocumentParseError("JATS XML input exceeds the configured extraction limit")
    if b"<!ENTITY" in data.upper():
        raise DocumentParseError("JATS XML entity declarations are not allowed")
    parser = etree.XMLParser(
        resolve_entities=False,
        no_network=True,
        load_dtd=False,
        recover=False,
        huge_tree=False,
        remove_comments=True,
    )
    try:
        root = etree.fromstring(data, parser=parser)
    except etree.XMLSyntaxError as exc:
        raise DocumentParseError("JATS XML is malformed") from exc
    articles = root.xpath("self::*[local-name()='article'] | .//*[local-name()='article']")
    if not articles:
        raise DocumentParseError("XML document does not contain a JATS article")

    output = TextAccumulator(max_chars)
    article_metadata: list[dict[str, Any]] = []
    for article_index, article in enumerate(articles, start=1):
        output.add(f"[[article:{article_index}]]")
        title = _first_xpath_text(article, "./*[local-name()='front']//*[local-name()='article-title']")
        if title:
            output.add("[[article-title]]")
            output.add(title)
        identifiers = {
            str(node.get("pub-id-type")): _normalized_xml_text(node)
            for node in article.xpath("./*[local-name()='front']//*[local-name()='article-id']")
            if node.get("pub-id-type") and _normalized_xml_text(node)
        }
        license_text = _first_xpath_text(article, "./*[local-name()='front']//*[local-name()='license']")
        license_references = [
            value
            for node in article.xpath("./*[local-name()='front']//*[local-name()='license']//*[@*]")
            for value in node.attrib.values()
            if isinstance(value, str) and value.startswith("https://creativecommons.org/licenses/")
        ]
        article_metadata.append(
            {
                "identifiers": identifiers,
                "license": license_text,
                "license_references": sorted(set(license_references)),
                "title": title,
            }
        )
        for abstract in article.xpath("./*[local-name()='front']//*[local-name()='abstract']"):
            output.add("[[abstract]]")
            output.add(_normalized_xml_text(abstract))
        for body in article.xpath("./*[local-name()='body']"):
            for child in body:
                local_name = etree.QName(child).localname
                if local_name == "sec":
                    _add_jats_section(output, child)
                elif local_name in {"p", "list", "table-wrap", "disp-quote", "boxed-text"}:
                    output.add(_normalized_xml_text(child))

    return ParsedDocument(
        output.render(),
        {"article_count": len(articles), "articles": article_metadata, "format": "JATS XML"},
        "lxml-jats",
        importlib.metadata.version("lxml"),
    )


def _read_text(path: Path, max_chars: int) -> tuple[str, str]:
    data = path.read_bytes()
    if len(data) > max_chars * 4:
        raise DocumentParseError("Text input exceeds the configured extraction limit")
    for encoding in ("utf-8-sig", "utf-16", "gb18030"):
        try:
            text = data.decode(encoding)
        except UnicodeDecodeError:
            continue
        if len(text) > max_chars:
            raise DocumentParseError(f"Extracted text exceeds the configured {max_chars} character limit")
        return text, encoding
    raise DocumentParseError("Text encoding is not UTF-8, UTF-16, or GB18030")


def _safe_office_member_name(value: str) -> str:
    if not value or "\x00" in value or "\\" in value or value.startswith("/"):
        raise DocumentParseError("Office document contains an unsafe archive member path")
    normalized = PurePosixPath(value)
    if any(part in {"", ".", ".."} for part in normalized.parts):
        raise DocumentParseError("Office document contains an unsafe archive member path")
    if normalized.parts and normalized.parts[0].endswith(":"):
        raise DocumentParseError("Office document contains an unsafe archive member path")
    return normalized.as_posix()


def _validate_office_archive(path: Path) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            members = archive.infolist()
            if len(members) > OFFICE_MAX_ARCHIVE_MEMBERS:
                raise DocumentParseError("Office document contains too many archive members")
            total = 0
            compressed = 0
            names: set[str] = set()
            for item in members:
                normalized_name = _safe_office_member_name(item.filename)
                identity = normalized_name.casefold()
                if identity in names:
                    raise DocumentParseError("Office document contains duplicate archive member names")
                names.add(identity)
                if item.flag_bits & 0x1:
                    raise DocumentParseError("Encrypted Office archive members require a decryption workflow")
                mode = (item.external_attr >> 16) & 0xFFFF
                if mode and stat.S_ISLNK(mode):
                    raise DocumentParseError("Office document contains a symbolic-link archive member")
                if item.compress_type not in {zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED}:
                    raise DocumentParseError("Office document uses an unsupported archive compression method")
                if item.is_dir():
                    continue
                if item.file_size < 0 or item.compress_size < 0:
                    raise DocumentParseError("Office document contains invalid archive member sizes")
                if item.file_size > OFFICE_MAX_MEMBER_UNCOMPRESSED_BYTES:
                    raise DocumentParseError("Office document contains an oversized archive member")
                member_compressed = max(item.compress_size, 1)
                if (
                    item.file_size > OFFICE_COMPRESSION_RATIO_MIN_BYTES
                    and item.file_size / member_compressed > OFFICE_MAX_COMPRESSION_RATIO
                ):
                    raise DocumentParseError("Office document has an unsafe archive member compression ratio")
                total += item.file_size
                compressed += member_compressed
                if total > OFFICE_MAX_UNCOMPRESSED_BYTES:
                    raise DocumentParseError("Office document expands beyond the configured safety limit")
    except (OSError, zipfile.BadZipFile, zipfile.LargeZipFile) as exc:
        raise DocumentParseError("Office document is not a valid Open XML archive") from exc
    if total > OFFICE_COMPRESSION_RATIO_MIN_BYTES and total / max(compressed, 1) > OFFICE_MAX_COMPRESSION_RATIO:
        raise DocumentParseError("Office document has an unsafe archive compression ratio")


PARSERS: dict[str, Parser] = {
    ".csv": _parse_delimited,
    ".docx": _parse_docx,
    ".htm": _parse_html,
    ".html": _parse_html,
    ".jpeg": _parse_ocr_image,
    ".jpg": _parse_ocr_image,
    ".json": _parse_json,
    ".md": _parse_text,
    ".nxml": _parse_jats_xml,
    ".pdf": _parse_pdf,
    ".png": _parse_ocr_image,
    ".pptx": _parse_pptx,
    ".txt": _parse_text,
    ".tif": _parse_ocr_image,
    ".tiff": _parse_ocr_image,
    ".xml": _parse_jats_xml,
    ".xlsx": _parse_xlsx,
}


def _register_scientific_parsers() -> None:
    if ".sdf" in PARSERS:
        return
    from pharma_intel.ingest.scientific_parsers import parse_macromolecular_structure, parse_mol, parse_sdf

    PARSERS.update(
        {
            ".cif": parse_macromolecular_structure,
            ".mmcif": parse_macromolecular_structure,
            ".mol": parse_mol,
            ".pdb": parse_macromolecular_structure,
            ".sdf": parse_sdf,
        }
    )
