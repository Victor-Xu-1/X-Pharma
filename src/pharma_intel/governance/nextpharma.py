from __future__ import annotations

import hashlib
import re
from collections import Counter
from dataclasses import dataclass
from datetime import UTC, date, datetime
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from pharma_intel.governance.schemas import (
    Citation,
    EntityReference,
    ProgramFact,
    ProgramMilestoneFact,
    ProgramOrganizationFact,
    ProgramStatusHistoryFact,
    ProgramTargetFact,
)
from pharma_intel.models import EntityType, ProgramTargetRole

ADAPTER_NAME = "nextpharma_xlsx"
ADAPTER_VERSION = "1.0.0"
RESULT_SHEET = "检索结果"
CONDITION_SHEET = "检索条件"
MAX_WORKBOOK_ROWS = 100_000
MAX_WORKBOOK_COLUMNS = 200
AUTHORIZED_SCOPE_PREFIX = "user-provided:pharmcube-"

REQUIRED_HEADERS = frozenset(
    {
        "数据层级",
        "药品名称",
        "研发机构",
        "研发机构(类型)",
        "靶点",
        "Modality",
        "作用机制",
        "适应症",
        "全球最高研发阶段",
        "更新日期",
        "全球研发状态",
        "全球最高研发阶段开始日期",
        "中国最高研发阶段",
        "中国最高研发阶段开始日期",
        "创新类型",
        "药品类别",
        "药品标签",
        "适应症领域",
        "NPUID",
        "研发机构(全)",
        "原研机构所在国家/地区",
        "原研机构权益地区",
    }
)

_TARGET_CANONICAL_NAMES = {
    "ifnα2": "IFNA2",
    "ifn-a2": "IFNA2",
    "ifn-α2": "IFNA2",
    "ifnγ": "IFNG",
    "ifn-gamma": "IFNG",
    "ifnγr1": "IFNGR1",
    "ifn-γr1": "IFNGR1",
    "ifnar-1": "IFNAR1",
    "ifnar1": "IFNAR1",
    "ifnar-2": "IFNAR2",
    "ifnar2": "IFNAR2",
    "il-22": "IL22",
    "il-22bp": "IL22RA2",
    "il-22rα1": "IL22RA1",
    "il-22ra1": "IL22RA1",
    "il-27": "IL27",
    "il-27rα": "IL27RA",
    "il-27ra": "IL27RA",
    "il-6r": "IL6R",
    "gp130": "IL6ST",
    "tnf-α": "TNF",
    "tnfα": "TNF",
    "tnfr1": "TNFRSF1A",
    "tnfr2": "TNFRSF1B",
}

_PLACEHOLDERS = frozenset({"", "-", "--", "n/a", "na", "none", "null"})
_TARGET_PLACEHOLDERS = _PLACEHOLDERS | frozenset({"not available", "unknown", "未披露", "暂无"})
_SPLIT_VALUES = re.compile(r"(?:[;；\r\n]+|\s+\+\s+)")
_QUERY_TARGET = re.compile(r"靶点\s*=\s*\(([^)]+)\)")


class NextPharmaWorkbookError(ValueError):
    pass


@dataclass(frozen=True)
class NextPharmaRecord:
    row_number: int
    source_locator: str
    source_quote: str
    fact: ProgramFact


@dataclass(frozen=True)
class NextPharmaExtraction:
    records: tuple[NextPharmaRecord, ...]
    query_target: str | None
    header_sha256: str
    rows_seen: int
    skipped_rows: dict[str, int]


def is_authorized_nextpharma_asset(
    *,
    file_name: str,
    extension: str,
    authorization_scopes: list[str],
) -> bool:
    return (
        extension.casefold() == ".xlsx"
        and file_name.casefold().startswith("nextpharma")
        and any(scope.casefold().startswith(AUTHORIZED_SCOPE_PREFIX) for scope in authorization_scopes)
    )


def adapter_policy_manifest() -> dict[str, object]:
    return {
        "adapter": ADAPTER_NAME,
        "version": ADAPTER_VERSION,
        "required_headers": sorted(REQUIRED_HEADERS),
        "accepted_data_level": "适应症信息",
        "max_rows": MAX_WORKBOOK_ROWS,
        "max_columns": MAX_WORKBOOK_COLUMNS,
        "authorization_scope_prefix": AUTHORIZED_SCOPE_PREFIX,
    }


def parse_nextpharma_workbook(content: bytes) -> NextPharmaExtraction:
    if not content:
        raise NextPharmaWorkbookError("NextPharma workbook is empty")
    try:
        workbook = load_workbook(
            BytesIO(content),
            read_only=True,
            data_only=True,
            keep_links=False,
        )
    except Exception as exc:
        raise NextPharmaWorkbookError("NextPharma workbook is not a readable XLSX archive") from exc

    try:
        if RESULT_SHEET not in workbook.sheetnames:
            raise NextPharmaWorkbookError(f"NextPharma workbook is missing the {RESULT_SHEET!r} worksheet")
        worksheet = workbook[RESULT_SHEET]
        if worksheet.max_row > MAX_WORKBOOK_ROWS + 1:
            raise NextPharmaWorkbookError("NextPharma workbook exceeds the deterministic row limit")
        if worksheet.max_column > MAX_WORKBOOK_COLUMNS:
            raise NextPharmaWorkbookError("NextPharma workbook exceeds the deterministic column limit")

        rows = worksheet.iter_rows(values_only=True)
        try:
            raw_headers = next(rows)
        except StopIteration as exc:
            raise NextPharmaWorkbookError("NextPharma result worksheet is empty") from exc
        headers = [_clean_text(value) or "" for value in raw_headers]
        duplicate_headers = [name for name, count in Counter(headers).items() if name and count > 1]
        if duplicate_headers:
            raise NextPharmaWorkbookError(
                "NextPharma result worksheet contains duplicate headers: " + ", ".join(sorted(duplicate_headers))
            )
        missing_headers = sorted(REQUIRED_HEADERS.difference(headers))
        if missing_headers:
            raise NextPharmaWorkbookError(
                "NextPharma result worksheet is missing required headers: " + ", ".join(missing_headers)
            )
        header_index = {name: index for index, name in enumerate(headers) if name}
        query_target = _workbook_query_target(workbook)
        records: list[NextPharmaRecord] = []
        skipped: Counter[str] = Counter()
        rows_seen = 0
        for row_number, raw_row in enumerate(rows, start=2):
            rows_seen += 1
            row = {name: raw_row[index] if index < len(raw_row) else None for name, index in header_index.items()}
            data_level = _clean_text(row["数据层级"])
            if data_level != "适应症信息":
                skipped["non_indication_row"] += 1
                continue
            missing = [
                field for field in ("药品名称", "靶点", "适应症", "全球最高研发阶段") if _clean_text(row[field]) is None
            ]
            if missing:
                skipped["missing_required_value"] += 1
                continue
            record = _program_record(row, row_number, query_target)
            records.append(record)
        if not records:
            raise NextPharmaWorkbookError("NextPharma workbook contains no publishable indication records")
        header_sha256 = hashlib.sha256("\u001f".join(headers).encode("utf-8")).hexdigest()
        return NextPharmaExtraction(
            records=tuple(records),
            query_target=query_target,
            header_sha256=header_sha256,
            rows_seen=rows_seen,
            skipped_rows=dict(sorted(skipped.items())),
        )
    finally:
        workbook.close()


def _program_record(row: dict[str, Any], row_number: int, query_target: str | None) -> NextPharmaRecord:
    drug_name = _required_text(row["药品名称"], "药品名称", row_number)
    indication_name = _required_text(row["适应症"], "适应症", row_number)
    phase = _required_text(row["全球最高研发阶段"], "全球最高研发阶段", row_number)
    source_target_names = _split_values(_required_text(row["靶点"], "靶点", row_number))
    if not source_target_names:
        raise NextPharmaWorkbookError(f"NextPharma row {row_number} has no target values")
    targets = _target_facts(source_target_names, query_target)
    organizations = _organization_facts(row)
    npuid = _clean_text(row.get("NPUID"))
    drug_external_ids = (
        {"pharmcube_npuid": npuid} if npuid is not None else {"pharmcube_drug": _scoped_identity("drug", drug_name)}
    )
    status = _clean_text(row.get("全球研发状态"))
    program_status = _program_status(status)
    global_phase = _clean_text(row.get("全球最高研发阶段"))
    china_phase = _clean_text(row.get("中国最高研发阶段"))
    status_date = _optional_datetime(row.get("更新日期"), "更新日期", row_number)
    global_phase_started_at = _optional_datetime(
        row.get("全球最高研发阶段开始日期"),
        "全球最高研发阶段开始日期",
        row_number,
    )
    china_phase_started_at = _optional_datetime(
        row.get("中国最高研发阶段开始日期"),
        "中国最高研发阶段开始日期",
        row_number,
    )
    quote_fields = (
        "数据层级",
        "药品名称",
        "研发机构",
        "靶点",
        "Modality",
        "作用机制",
        "适应症",
        "全球最高研发阶段",
        "更新日期",
        "全球研发状态",
        "NPUID",
    )
    source_quote = "\t".join(_cell_text(row.get(field)) for field in quote_fields)
    source_quote = source_quote[:4000]
    source_locator = f"sheet:{RESULT_SHEET}/row:{row_number}"
    fact = ProgramFact(
        fact_kind="program",
        drug=EntityReference(
            entity_type=EntityType.DRUG,
            name=drug_name,
            external_ids=drug_external_ids,
        ),
        targets=targets,
        indication=EntityReference(
            entity_type=EntityType.DISEASE,
            name=indication_name,
            external_ids={"pharmcube_indication": _scoped_identity("indication", indication_name)},
        ),
        organizations=organizations,
        phase=phase,
        status=status,
        status_date=status_date,
        modality=_bounded_text(row.get("Modality"), 160, "Modality", row_number),
        innovation_type=_bounded_text(row.get("创新类型"), 120, "创新类型", row_number),
        therapeutic_area=_first_value(row.get("适应症领域"), 120, "适应症领域", row_number),
        drug_category=_bounded_text(row.get("药品类别"), 120, "药品类别", row_number),
        mechanism_of_action=_bounded_text(row.get("作用机制"), 500, "作用机制", row_number),
        geography="全球",
        global_phase=global_phase,
        china_phase=china_phase,
        global_phase_started_at=global_phase_started_at,
        china_phase_started_at=china_phase_started_at,
        development_rights_regions=_split_values(row.get("原研机构权益地区")),
        program_tags=_split_values(row.get("药品标签")),
        status_history=_status_history(row, status, program_status, row_number),
        milestones=_milestones(row, row_number),
        citation=Citation(locator=source_locator, quote=source_quote, confidence=1.0),
    )
    return NextPharmaRecord(
        row_number=row_number,
        source_locator=source_locator,
        source_quote=source_quote,
        fact=fact,
    )


def _target_facts(source_names: list[str], query_target: str | None) -> list[ProgramTargetFact]:
    target_names = list(
        dict.fromkeys(_canonical_target_name(name) for name in source_names if not _is_target_placeholder(name))
    )
    if not target_names:
        raise NextPharmaWorkbookError("NextPharma row has no meaningful target values")
    canonical_query = (
        _canonical_target_name(query_target)
        if query_target is not None and not _is_target_placeholder(query_target)
        else None
    )
    if canonical_query in target_names:
        target_names.remove(canonical_query)
        target_names.insert(0, canonical_query)
    return [
        ProgramTargetFact(
            role=ProgramTargetRole.PRIMARY if index == 0 else ProgramTargetRole.COMBINATION,
            entity=EntityReference(
                entity_type=EntityType.TARGET,
                name=name,
                external_ids={"pharmcube_target": _scoped_identity("target", name)},
            ),
        )
        for index, name in enumerate(target_names)
    ]


def _is_target_placeholder(value: str) -> bool:
    return " ".join(value.casefold().split()) in _TARGET_PLACEHOLDERS


def _organization_facts(row: dict[str, Any]) -> list[ProgramOrganizationFact]:
    names = _split_values(row.get("研发机构"))
    full_names = _split_values(row.get("研发机构(全)"))
    descriptors = _split_values(row.get("研发机构(类型)"))
    if full_names and len(full_names) == len(names):
        names = full_names
    if not names:
        return []
    originator_index = next(
        (index for index, descriptor in enumerate(descriptors[: len(names)]) if "原研" in descriptor),
        0,
    )
    if len(full_names) == 1 and 0 <= originator_index < len(names):
        names[originator_index] = full_names[0]
    country = _clean_text(row.get("原研机构所在国家/地区"))
    facts: list[ProgramOrganizationFact] = []
    seen: set[str] = set()
    for index, name in enumerate(names):
        normalized = " ".join(name.casefold().split())
        if normalized in seen:
            continue
        seen.add(normalized)
        role = "originator" if index == originator_index else "collaborator"
        descriptor = descriptors[index] if index < len(descriptors) else None
        facts.append(
            ProgramOrganizationFact(
                role=role,
                entity=EntityReference(
                    entity_type=EntityType.ORGANIZATION,
                    name=name,
                    external_ids={"pharmcube_organization": _scoped_identity("organization", name)},
                ),
                country_region=country if role == "originator" else None,
                organization_type=descriptor[:120] if descriptor else None,
            )
        )
    facts.sort(key=lambda fact: fact.role != "originator")
    return facts


def _status_history(
    row: dict[str, Any],
    status: str | None,
    program_status: str | None,
    row_number: int,
) -> list[ProgramStatusHistoryFact]:
    definitions = (
        ("全球", "全球最高研发阶段", "全球最高研发阶段开始日期"),
        ("中国", "中国最高研发阶段", "中国最高研发阶段开始日期"),
        ("美国", "美国最高研发阶段", "美国最高研发阶段开始日期"),
        ("欧洲", "欧洲最高研发阶段", "欧洲最高研发阶段开始日期"),
        ("日本", "日本最高研发阶段", "日本最高研发阶段开始日期"),
        ("其他", "其他最高研发阶段", "其他最高研发阶段开始日期"),
    )
    history: list[ProgramStatusHistoryFact] = []
    for geography, phase_field, date_field in definitions:
        phase = _clean_text(row.get(phase_field))
        effective_at = _optional_datetime(row.get(date_field), date_field, row_number)
        if phase is None or effective_at is None:
            continue
        history.append(
            ProgramStatusHistoryFact(
                phase=phase,
                status=status,
                program_status=program_status,
                effective_at=effective_at,
                geography=geography,
            )
        )
    return history


def _milestones(row: dict[str, Any], row_number: int) -> list[ProgramMilestoneFact]:
    definitions = (
        ("first_global_trial", "全球首次临床登记日期", "全球首次临床登记"),
        ("china_ind", "中国首个IND日期", "中国首次 IND"),
        ("china_nda", "中国首个NDA日期", "中国首次 NDA"),
        ("development_discontinued", "停研日期", "研发停止"),
    )
    milestones: list[ProgramMilestoneFact] = []
    for milestone_type, date_field, title in definitions:
        occurred_at = _optional_datetime(row.get(date_field), date_field, row_number)
        if occurred_at is None:
            continue
        milestones.append(
            ProgramMilestoneFact(
                milestone_type=milestone_type,
                title=title,
                occurred_at=occurred_at,
                geography="中国" if milestone_type.startswith("china_") else "全球",
            )
        )
    return milestones


def _workbook_query_target(workbook: Any) -> str | None:
    if CONDITION_SHEET not in workbook.sheetnames:
        return None
    for row in workbook[CONDITION_SHEET].iter_rows(values_only=True):
        for value in row:
            text = _clean_text(value)
            if text is None:
                continue
            match = _QUERY_TARGET.search(text)
            if match:
                return match.group(1).strip()
    return None


def _canonical_target_name(value: str) -> str:
    compact = re.sub(r"\s+", "", value).casefold()
    return _TARGET_CANONICAL_NAMES.get(compact, value.strip())


def _scoped_identity(kind: str, value: str) -> str:
    normalized = " ".join(value.casefold().split())
    return hashlib.sha256(f"{kind}\u001f{normalized}".encode()).hexdigest()


def _program_status(value: str | None) -> str | None:
    if value is None:
        return None
    normalized = value.casefold()
    if normalized == "active":
        return "active"
    if normalized in {"inactive", "discontinued", "terminated", "withdrawn"}:
        return "inactive"
    return "unknown"


def _optional_datetime(value: Any, field_name: str, row_number: int) -> datetime | None:
    text = _clean_text(value)
    if text is None:
        return None
    if isinstance(value, datetime):
        return value if value.tzinfo else value.replace(tzinfo=UTC)
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day, tzinfo=UTC)
    try:
        parsed = datetime.fromisoformat(text.replace("Z", "+00:00"))
    except ValueError as exc:
        raise NextPharmaWorkbookError(
            f"NextPharma row {row_number} has an invalid {field_name} value: {text!r}"
        ) from exc
    return parsed if parsed.tzinfo else parsed.replace(tzinfo=UTC)


def _bounded_text(value: Any, max_length: int, field_name: str, row_number: int) -> str | None:
    text = _clean_text(value)
    if text is not None and len(text) > max_length:
        raise NextPharmaWorkbookError(
            f"NextPharma row {row_number} has an overlong {field_name} value ({len(text)} characters)"
        )
    return text


def _first_value(value: Any, max_length: int, field_name: str, row_number: int) -> str | None:
    values = _split_values(value)
    return _bounded_text(values[0], max_length, field_name, row_number) if values else None


def _required_text(value: Any, field_name: str, row_number: int) -> str:
    text = _clean_text(value)
    if text is None:
        raise NextPharmaWorkbookError(f"NextPharma row {row_number} is missing {field_name}")
    return text


def _split_values(value: Any) -> list[str]:
    text = _clean_text(value)
    if text is None:
        return []
    return list(dict.fromkeys(item.strip() for item in _SPLIT_VALUES.split(text) if _clean_text(item) is not None))


def _clean_text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    if text.casefold() in _PLACEHOLDERS:
        return None
    return text


def _cell_text(value: Any) -> str:
    return _clean_text(value) or ""
