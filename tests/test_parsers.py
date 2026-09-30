from __future__ import annotations

import stat
import zipfile
from importlib import import_module
from io import BytesIO
from pathlib import Path
from typing import Any

import gemmi
import pytest
from bs4 import BeautifulSoup
from docx import Document
from openpyxl import Workbook
from PIL import Image
from pptx import Presentation
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen.canvas import Canvas

from pharma_intel.ingest import parsers, scientific_parsers
from pharma_intel.ingest.parsers import DocumentOcrRequired, DocumentParseError, parse_document

Chem: Any = import_module("rdkit.Chem")


def _office_archive(path: Path, members: list[tuple[zipfile.ZipInfo | str, bytes]]) -> None:
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, payload in members:
            archive.writestr(name, payload)


def _mark_zip_members_encrypted(path: Path) -> None:
    payload = bytearray(path.read_bytes())
    for signature, flag_offset in ((b"PK\x03\x04", 6), (b"PK\x01\x02", 8)):
        cursor = 0
        while (cursor := payload.find(signature, cursor)) >= 0:
            flags = int.from_bytes(payload[cursor + flag_offset : cursor + flag_offset + 2], "little")
            payload[cursor + flag_offset : cursor + flag_offset + 2] = (flags | 1).to_bytes(2, "little")
            cursor += len(signature)
    path.write_bytes(payload)


def test_regular_document_does_not_load_scientific_native_extensions(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    source = tmp_path / "evidence.md"
    source.write_text("EGFR evidence", encoding="utf-8")

    def unexpected_registration() -> None:
        raise AssertionError("regular documents must not load scientific parsers")

    monkeypatch.setattr(parsers, "_register_scientific_parsers", unexpected_registration)

    assert parse_document(source, 100_000).text == "EGFR evidence"


def test_json_parser_validates_and_canonicalizes_source_records(tmp_path: Path) -> None:
    source = tmp_path / "NCT00000001.json"
    source.write_text('{"z": 1, "trial": {"nctId": "NCT00000001"}}', encoding="utf-8")

    parsed = parse_document(source, 100_000)

    assert parsed.text == '{"trial":{"nctId":"NCT00000001"},"z":1}'
    assert parsed.metadata == {"encoding": "utf-8-sig", "top_level_type": "dict"}
    assert parsed.parser_name == "python-json"


def test_json_parser_rejects_invalid_json(tmp_path: Path) -> None:
    source = tmp_path / "invalid.json"
    source.write_text('{"trial":', encoding="utf-8")

    with pytest.raises(DocumentParseError, match="JSON document is not syntactically valid"):
        parse_document(source, 100_000)


@pytest.mark.parametrize(
    ("members", "message"),
    [
        ([("../escape.xml", b"unsafe")], "unsafe archive member path"),
        ([("word/document.xml", b"one"), ("WORD/document.xml", b"two")], "duplicate archive member"),
    ],
)
def test_office_archive_preflight_rejects_unsafe_names(
    tmp_path: Path,
    members: list[tuple[zipfile.ZipInfo | str, bytes]],
    message: str,
) -> None:
    source = tmp_path / "unsafe.docx"
    _office_archive(source, members)

    with pytest.raises(DocumentParseError, match=message):
        parsers._validate_office_archive(source)


def test_office_archive_preflight_rejects_symlinks_expansion_and_member_fanout(tmp_path: Path) -> None:
    symlink = zipfile.ZipInfo("word/link.xml")
    symlink.create_system = 3
    symlink.external_attr = (stat.S_IFLNK | 0o777) << 16
    symlink_path = tmp_path / "symlink.docx"
    _office_archive(symlink_path, [(symlink, b"word/document.xml")])
    with pytest.raises(DocumentParseError, match="symbolic-link"):
        parsers._validate_office_archive(symlink_path)

    expansion = tmp_path / "expansion.docx"
    _office_archive(expansion, [("word/document.xml", b"A" * (parsers.OFFICE_COMPRESSION_RATIO_MIN_BYTES + 1))])
    with pytest.raises(DocumentParseError, match="member compression ratio"):
        parsers._validate_office_archive(expansion)

    fanout = tmp_path / "fanout.xlsx"
    _office_archive(
        fanout,
        [(f"xl/worksheets/sheet{index}.xml", b"") for index in range(parsers.OFFICE_MAX_ARCHIVE_MEMBERS + 1)],
    )
    with pytest.raises(DocumentParseError, match="too many archive members"):
        parsers._validate_office_archive(fanout)


def test_office_archive_preflight_rejects_encrypted_members(tmp_path: Path) -> None:
    encrypted = tmp_path / "encrypted.docx"
    _office_archive(encrypted, [("word/document.xml", b"encrypted")])
    _mark_zip_members_encrypted(encrypted)

    with zipfile.ZipFile(encrypted) as archive:
        assert archive.infolist()[0].flag_bits & 1
    with pytest.raises(DocumentParseError, match="Encrypted Office archive members"):
        parsers._validate_office_archive(encrypted)


def test_real_document_parsers_extract_source_text(tmp_path: Path) -> None:
    pdf = tmp_path / "paper.pdf"
    canvas = Canvas(str(pdf))
    canvas.drawString(72, 720, "EGFR inhibitor evidence")
    canvas.save()

    docx = tmp_path / "report.docx"
    document = Document()
    document.add_paragraph("Clinical program summary")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Target"
    table.cell(0, 1).text = "EGFR"
    document.save(str(docx))

    pptx = tmp_path / "deck.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Competitive landscape"
    slide.placeholders[1].text = "Phase 2 asset"
    presentation.save(str(pptx))

    xlsx = tmp_path / "activities.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "Activities"
    worksheet.append(["compound", "IC50", "unit"])
    worksheet.append(["CMP-1", 12.5, "nM"])
    workbook.save(xlsx)

    parsed_pdf = parse_document(pdf, 100_000)
    parsed_docx = parse_document(docx, 100_000)
    parsed_pptx = parse_document(pptx, 100_000)
    parsed_xlsx = parse_document(xlsx, 100_000)

    assert "EGFR inhibitor evidence" in parsed_pdf.text
    assert "[[page:1]]" in parsed_pdf.text
    assert "Clinical program summary" in parsed_docx.text
    assert "Target\tEGFR" in parsed_docx.text
    assert "[[slide:1]]" in parsed_pptx.text
    assert "Phase 2 asset" in parsed_pptx.text
    assert "[[sheet:Activities]]" in parsed_xlsx.text
    assert "CMP-1\t12.5\tnM" in parsed_xlsx.text


def test_parser_rejects_fake_or_empty_content(tmp_path: Path) -> None:
    fake_pdf = tmp_path / "fake.pdf"
    fake_pdf.write_text("not a pdf", encoding="utf-8")

    with pytest.raises(DocumentParseError):
        parse_document(fake_pdf, 100_000)


def test_parser_routes_document_images_and_scan_only_pdf_to_ocr(tmp_path: Path) -> None:
    image = Image.new("RGB", (320, 120), "white")
    png = tmp_path / "scan.png"
    image.save(png)
    pdf = tmp_path / "scan.pdf"
    payload = BytesIO()
    image.save(payload, format="PNG")
    canvas = Canvas(str(pdf))
    canvas.drawImage(ImageReader(payload), 72, 600, width=320, height=120)
    canvas.save()

    with pytest.raises(DocumentOcrRequired, match="requires OCR"):
        parse_document(png, 100_000)
    with pytest.raises(DocumentOcrRequired, match="requires OCR"):
        parse_document(pdf, 100_000)


def test_html_parser_uses_lxml_for_untrusted_input(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    html = tmp_path / "evidence.html"
    html.write_text(
        "<!doctype html><html><head><title>EGFR evidence</title><script>discard()</script></head>"
        "<body><p>Osimertinib IC50 12 nM</p></body></html>",
        encoding="utf-8",
    )
    original = BeautifulSoup
    selected_features: list[str | None] = []

    def capture_features(markup: str, features: str | None = None) -> object:
        selected_features.append(features)
        return original(markup, features)

    monkeypatch.setattr("pharma_intel.ingest.parsers.BeautifulSoup", capture_features)

    parsed = parse_document(html, 100_000)

    assert selected_features == ["lxml"]
    assert parsed.parser_name == "beautifulsoup4+lxml"
    assert parsed.metadata["title"] == "EGFR evidence"
    assert "Osimertinib IC50 12 nM" in parsed.text
    assert "discard" not in parsed.text


def test_jats_xml_parser_extracts_governed_article_text_and_metadata(tmp_path: Path) -> None:
    source = tmp_path / "article.xml"
    source.write_text(
        """<?xml version="1.0"?>
<!DOCTYPE article PUBLIC "-//NLM//DTD JATS (Z39.96) Journal Archiving and Interchange DTD v1.4 20241031//EN" "https://dtd.nlm.nih.gov/archiving/1.4/JATS-archivearticle1-4.dtd">
<article xmlns:xlink="http://www.w3.org/1999/xlink">
  <front><article-meta>
    <article-id pub-id-type="pmcid">PMC123</article-id>
    <article-id pub-id-type="doi">10.1000/example</article-id>
    <title-group><article-title>EGFR combination evidence</article-title></title-group>
    <permissions><license><license-p>
      <ext-link xlink:href="https://creativecommons.org/licenses/by/4.0/">CC BY 4.0</ext-link>
    </license-p></license></permissions>
    <abstract><p>Osimertinib targets mutant EGFR.</p></abstract>
  </article-meta></front>
  <body><sec id="results"><title>Results</title>
    <p>Volasertib sensitized EGFR-mutant cells.</p>
    <sec><title>Activity</title><p>Cell viability was measured after 72 h.</p></sec>
  </sec></body>
</article>""",
        encoding="utf-8",
    )

    parsed = parse_document(source, 100_000)

    assert parsed.parser_name == "lxml-jats"
    assert parsed.metadata["format"] == "JATS XML"
    assert parsed.metadata["articles"][0]["identifiers"] == {
        "doi": "10.1000/example",
        "pmcid": "PMC123",
    }
    assert parsed.metadata["articles"][0]["license_references"] == ["https://creativecommons.org/licenses/by/4.0/"]
    assert "[[article-title]]\nEGFR combination evidence" in parsed.text
    assert "[[section:Results]]" in parsed.text
    assert "[[section:Activity]]" in parsed.text
    assert "Cell viability was measured after 72 h." in parsed.text


def test_jats_xml_parser_rejects_entities_malformed_xml_and_non_jats(tmp_path: Path) -> None:
    entity = tmp_path / "entity.xml"
    entity.write_text(
        '<?xml version="1.0"?><!DOCTYPE article [<!ENTITY xxe SYSTEM "file:///etc/passwd">]>'
        "<article><body><p>&xxe;</p></body></article>",
        encoding="utf-8",
    )
    malformed = tmp_path / "malformed.nxml"
    malformed.write_text("<article><body>", encoding="utf-8")
    generic = tmp_path / "generic.xml"
    generic.write_text("<configuration><secret>value</secret></configuration>", encoding="utf-8")

    with pytest.raises(DocumentParseError, match="entity declarations"):
        parse_document(entity, 100_000)
    with pytest.raises(DocumentParseError, match="malformed"):
        parse_document(malformed, 100_000)
    with pytest.raises(DocumentParseError, match="JATS article"):
        parse_document(generic, 100_000)


def test_rdkit_parsers_extract_and_standardize_real_sdf_and_mol(tmp_path: Path) -> None:
    sdf = tmp_path / "compounds.sdf"
    writer = Chem.SDWriter(str(sdf))
    for name, smiles in (("Ethanol", "CCO"), ("Acetic acid", "CC(=O)O")):
        molecule = Chem.MolFromSmiles(smiles)
        assert molecule is not None
        molecule.SetProp("_Name", name)
        molecule.SetProp("source_batch", "BATCH-42")
        writer.write(molecule)
    writer.close()

    mol = tmp_path / "ethanol.mol"
    molecule = Chem.MolFromSmiles("CCO")
    assert molecule is not None
    molecule.SetProp("_Name", "Ethanol")
    Chem.MolToMolFile(molecule, str(mol))

    parsed_sdf = parse_document(sdf, 100_000)
    parsed_mol = parse_document(mol, 100_000)

    assert parsed_sdf.parser_name == "rdkit-sdf"
    assert parsed_sdf.metadata["scientific_format"] == "sdf"
    assert parsed_sdf.metadata["record_count"] == 2
    assert parsed_sdf.metadata["record_preview"][0]["standard_inchi_key"] == "LFQSCWFLJHTTHZ-UHFFFAOYSA-N"
    assert "[[molecule:2]]" in parsed_sdf.text
    assert "property.source_batch: BATCH-42" in parsed_sdf.text
    assert parsed_mol.parser_name == "rdkit-mol"
    assert parsed_mol.metadata["record_count"] == 1
    assert "molecular_formula: C2H6O" in parsed_mol.text


def test_gemmi_parsers_extract_real_pdb_and_mmcif_metadata(tmp_path: Path) -> None:
    pdb = tmp_path / "complex.pdb"
    pdb.write_text(
        "HEADER    TEST STRUCTURE\n"
        "ATOM      1  N   GLY A   1      11.104  13.207  10.000  1.00 20.00           N\n"
        "ATOM      2  CA  GLY A   1      12.000  12.100  10.000  1.00 20.00           C\n"
        "HETATM    3  C1  LIG A 101      14.000  11.000  10.000  1.00 20.00           C\n"
        "END\n",
        encoding="ascii",
    )
    structure = gemmi.read_structure(str(pdb))
    mmcif = tmp_path / "complex.mmcif"
    structure.make_mmcif_document().write_file(str(mmcif))

    parsed_pdb = parse_document(pdb, 100_000)
    parsed_mmcif = parse_document(mmcif, 100_000)

    for parsed, expected_format in ((parsed_pdb, "pdb"), (parsed_mmcif, "mmcif")):
        assert parsed.parser_name == f"gemmi-{expected_format}"
        assert parsed.metadata["scientific_format"] == expected_format
        assert parsed.metadata["model_count"] == 1
        assert parsed.metadata["chain_ids"] == ["A"]
        assert parsed.metadata["residue_count"] == 2
        assert parsed.metadata["atom_count"] == 3
        assert parsed.metadata["ligands"] == [{"component_id": "LIG", "instance_count": 1}]
        assert "ligand: LIG instances=1" in parsed.text


@pytest.mark.parametrize("suffix", [".sdf", ".mol", ".pdb", ".cif", ".mmcif"])
def test_scientific_parsers_reject_corrupt_files(tmp_path: Path, suffix: str) -> None:
    source = tmp_path / f"corrupt{suffix}"
    source.write_bytes(b"not a valid scientific structure\x00")

    with pytest.raises(DocumentParseError):
        parse_document(source, 100_000)


def test_scientific_parsers_enforce_record_and_atom_limits(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    sdf = tmp_path / "too-many.sdf"
    writer = Chem.SDWriter(str(sdf))
    molecule = Chem.MolFromSmiles("CCO")
    assert molecule is not None
    writer.write(molecule)
    writer.write(molecule)
    writer.close()
    monkeypatch.setattr(scientific_parsers, "MAX_CHEMICAL_RECORDS", 1)
    with pytest.raises(DocumentParseError, match="molecule safety limit"):
        parse_document(sdf, 100_000)

    pdb = tmp_path / "too-many-atoms.pdb"
    pdb.write_text(
        "ATOM      1  N   GLY A   1      11.104  13.207  10.000  1.00 20.00           N\n"
        "ATOM      2  CA  GLY A   1      12.000  12.100  10.000  1.00 20.00           C\n"
        "END\n",
        encoding="ascii",
    )
    monkeypatch.setattr(scientific_parsers, "MAX_STRUCTURE_ATOMS", 1)
    with pytest.raises(DocumentParseError, match="atom safety limit"):
        parse_document(pdb, 100_000)
