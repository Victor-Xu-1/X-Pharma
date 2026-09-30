from __future__ import annotations

import hashlib
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import Workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from pharma_intel.config import Settings
from pharma_intel.governance.model_gateway import ExtractionResponse, OpenAICompatibleExtractionGateway
from pharma_intel.governance.nextpharma import NextPharmaWorkbookError, parse_nextpharma_workbook
from pharma_intel.governance.schemas import ExtractionEnvelope
from pharma_intel.governance.service import GovernanceService, _normalize_phase
from pharma_intel.models import (
    DataSource,
    DataSourceType,
    DevelopmentPhase,
    DevelopmentProgram,
    DevelopmentProgramOrganization,
    DevelopmentProgramTarget,
    Entity,
    EntityType,
    ExtractionRun,
    GovernanceStatus,
    SourceAsset,
    SourceDocument,
    SourceVersion,
    SourceVersionState,
    StagedFact,
    StageStatus,
    Tenant,
)
from pharma_intel.object_store import FileSystemObjectStore

HEADERS = [
    "数据层级",
    "药品名称",
    "研发机构",
    "研发机构(类型)",
    "研发机构(母公司)",
    "靶点",
    "Modality",
    "作用机制",
    "适应症",
    "全球最高研发阶段",
    "更新日期",
    "全球研发状态",
    "全球最高研发阶段开始日期",
    "中国最高研发阶段",
    "中国最高研发阶段开始日期",
    "临床结果数量",
    "创新类型",
    "靶点组合",
    "药品异名",
    "化学结构",
    "药品类别",
    "药品标签",
    "靶点类型",
    "Modality层级",
    "递送系统",
    "适应症领域",
    "剂型 (中国)",
    "更新类型",
    "DAR值",
    "NPUID",
    "增补数据",
    "研发机构(全)",
    "原研机构所在国家/地区",
    "原研机构权益地区",
    "中国研发赛道排名",
    "全球研发赛道排名",
    "美国最高研发阶段",
    "美国最高研发阶段开始日期",
    "欧洲最高研发阶段",
    "欧洲最高研发阶段开始日期",
    "日本最高研发阶段",
    "日本最高研发阶段开始日期",
    "其他最高研发阶段",
    "其他最高研发阶段开始日期",
    "分子最高研发阶段 (含仿制药)",
    "停研日期",
    "审评审批类型",
    "全球首次临床登记日期",
    "中国首个IND日期",
    "中国首个NDA日期",
    "临床结果最优评价",
    "交易数量",
    "交易总金额(百万)/美元",
    "交易首付款(百万)/美元",
    "中国化合物/序列专利到期日期",
    "美国化合物/序列专利到期日期",
]


class CountingGateway(OpenAICompatibleExtractionGateway):
    def __init__(self) -> None:
        self.calls = 0

    def extract(
        self,
        document_segment: str,
        *,
        fact_kind_allowlist: frozenset[str] | None = None,
        max_facts: int | None = None,
        source_profile: str | None = None,
    ) -> ExtractionResponse:
        del document_segment, fact_kind_allowlist, max_facts, source_profile
        self.calls += 1
        return ExtractionResponse(
            ExtractionEnvelope(document_type="spreadsheet", document_summary="Generic extraction", facts=[]),
            10,
            5,
        )


def test_nextpharma_parser_emits_indication_level_programs_with_row_provenance() -> None:
    workbook = _workbook_bytes(
        [
            _row(data_level="药品信息", indication="疾病甲;疾病乙"),
            _row(indication="疾病甲", phase="II期临床"),
            _row(indication="疾病乙", phase="I期临床", status="Inactive"),
        ],
        query_target="IL-27Rα",
    )

    extraction = parse_nextpharma_workbook(workbook)

    assert extraction.rows_seen == 3
    assert extraction.skipped_rows == {"non_indication_row": 1}
    assert extraction.query_target == "IL-27Rα"
    assert len(extraction.records) == 2
    first = extraction.records[0]
    assert first.row_number == 3
    assert first.source_locator == "sheet:检索结果/row:3"
    assert first.fact.drug.external_ids == {"pharmcube_npuid": "DR006303"}
    assert [item.entity.name for item in first.fact.targets] == ["IL27RA", "IL6ST"]
    assert [item.role.value for item in first.fact.targets] == ["primary", "combination"]
    assert first.fact.organizations[0].role == "originator"
    assert first.fact.organizations[0].entity.name == "Surface Oncology Inc."
    assert first.fact.phase == "II期临床"
    assert first.fact.citation.confidence == 1
    assert "casdozokitug" in first.source_quote


def test_nextpharma_parser_ignores_target_placeholders_without_losing_the_real_target() -> None:
    workbook = _workbook_bytes(
        [_row(targets="not available;IFN-α2")],
        query_target="IFNA2",
    )

    extraction = parse_nextpharma_workbook(workbook)

    assert [item.entity.name for item in extraction.records[0].fact.targets] == ["IFNA2"]
    assert [item.role.value for item in extraction.records[0].fact.targets] == ["primary"]


def test_nextpharma_parser_rejects_schema_drift_before_emitting_records() -> None:
    workbook = _workbook_bytes(
        [_row()],
        query_target="IL-27",
        headers=[header for header in HEADERS if header != "适应症"],
    )

    with pytest.raises(NextPharmaWorkbookError, match="missing required headers"):
        parse_nextpharma_workbook(workbook)


def test_authorized_nextpharma_governance_bypasses_llm_and_publishes_idempotent_programs(
    session: Session,
    tenant: Tenant,
    tmp_path: Path,
) -> None:
    workbook = _workbook_bytes(
        [
            _row(data_level="药品信息", indication="疾病甲;疾病乙"),
            _row(indication="疾病甲", phase="II期临床"),
            _row(indication="疾病乙", phase="I期临床"),
        ],
        query_target="IL-27Rα",
    )
    store, version = _version(
        session,
        tenant,
        tmp_path,
        file_name="NextPharma-基础查询-药品数据-20260716-test.xlsx",
        workbook=workbook,
        authorization_scopes=["user-provided:pharmcube-20260716"],
    )
    gateway = CountingGateway()
    service = GovernanceService(
        session,
        Settings(
            ai_governance_enabled=True,
            ai_base_url="https://model.test",
            ai_api_key="test-key",
            ai_model="must-not-be-called",
            object_store_root=tmp_path / "objects",
        ),
        store,
        tenant.id,
        gateway=gateway,
    )

    first = service.govern_version(version.id)
    repeated = service.govern_version(version.id)

    assert gateway.calls == 0
    assert first["run_id"] == repeated["run_id"]
    assert first["published"] == 2
    assert first["review_pending"] == first["rejected"] == first["conflict"] == 0
    assert session.scalar(select(func.count()).select_from(DevelopmentProgram)) == 2
    assert session.scalar(select(func.count()).select_from(DevelopmentProgramTarget)) == 4
    assert session.scalar(select(func.count()).select_from(DevelopmentProgramOrganization)) == 2
    target_names = set(session.scalars(select(Entity.name).where(Entity.entity_type == EntityType.TARGET)))
    assert target_names == {"IL27RA", "IL6ST"}
    phases = set(session.scalars(select(DevelopmentProgram.phase)))
    assert phases == {DevelopmentPhase.PHASE_1, DevelopmentPhase.PHASE_2}
    staged = list(session.scalars(select(StagedFact).order_by(StagedFact.source_locator)))
    assert [item.status for item in staged] == [GovernanceStatus.PUBLISHED, GovernanceStatus.PUBLISHED]
    assert [item.source_locator for item in staged] == ["sheet:检索结果/row:3", "sheet:检索结果/row:4"]
    run = session.get(ExtractionRun, first["run_id"])
    assert run is not None
    assert run.model_provider == "deterministic-adapter"
    assert run.model_name == "nextpharma_xlsx:1.0.0"
    assert run.input_tokens is None and run.output_tokens is None
    assert float(run.estimated_cost or 0) == 0
    assert run.structured_output is not None
    assert run.structured_output["deterministic_adapter"]["rows_emitted"] == 2
    assert version.state == SourceVersionState.PUBLISHED


def test_nextpharma_filename_without_authorization_scope_stays_on_generic_governance(
    session: Session,
    tenant: Tenant,
    tmp_path: Path,
) -> None:
    workbook = _workbook_bytes([_row()], query_target="IL-27")
    store, version = _version(
        session,
        tenant,
        tmp_path,
        file_name="NextPharma-untrusted.xlsx",
        workbook=workbook,
        authorization_scopes=["contract:unrelated-source"],
    )
    gateway = CountingGateway()
    result = GovernanceService(
        session,
        Settings(
            ai_governance_enabled=True,
            ai_base_url="https://model.test",
            ai_api_key="test-key",
            ai_model="generic-model",
            object_store_root=tmp_path / "objects",
        ),
        store,
        tenant.id,
        gateway=gateway,
    ).govern_version(version.id)

    assert gateway.calls == 1
    assert result["fact_count"] == 0


@pytest.mark.parametrize(
    ("source_phase", "expected"),
    [
        ("临床前", DevelopmentPhase.PRECLINICAL),
        ("申报临床", DevelopmentPhase.IND),
        ("I期临床", DevelopmentPhase.PHASE_1),
        ("I/II期临床", DevelopmentPhase.PHASE_1_2),
        ("II期临床", DevelopmentPhase.PHASE_2),
        ("II/III期临床", DevelopmentPhase.PHASE_2_3),
        ("III期临床", DevelopmentPhase.PHASE_3),
        ("申请上市", DevelopmentPhase.FILED),
        ("批准上市", DevelopmentPhase.APPROVED),
    ],
)
def test_nextpharma_chinese_phase_vocabulary(source_phase: str, expected: DevelopmentPhase) -> None:
    assert _normalize_phase(source_phase) == expected


def _row(
    *,
    data_level: str = "适应症信息",
    indication: str = "疾病甲",
    phase: str = "II期临床",
    status: str = "Active",
    targets: str = "gp130;IL-27Rα",
) -> dict[str, object]:
    return {
        "数据层级": data_level,
        "药品名称": "casdozokitug",
        "研发机构": "Surface Oncology",
        "研发机构(类型)": "Surface Oncology(原研)",
        "研发机构(母公司)": "Surface Oncology",
        "靶点": targets,
        "Modality": "单抗",
        "作用机制": "anti-IL-27单抗",
        "适应症": indication,
        "全球最高研发阶段": phase,
        "更新日期": "2025-05-27",
        "全球研发状态": status,
        "全球最高研发阶段开始日期": "2022-05-04",
        "中国最高研发阶段": "--",
        "中国最高研发阶段开始日期": "",
        "临床结果数量": 1,
        "创新类型": "创新药",
        "药品类别": "生物制品",
        "药品标签": "Potential First-in-Class;肿瘤免疫",
        "适应症领域": "肿瘤领域",
        "NPUID": "DR006303",
        "研发机构(全)": "Surface Oncology Inc.",
        "原研机构所在国家/地区": "美国",
        "原研机构权益地区": "全球;美国;中国(内地)",
        "美国最高研发阶段": "I期临床",
        "美国最高研发阶段开始日期": "2020-05-05",
        "全球首次临床登记日期": "2020-05-05",
    }


def _workbook_bytes(
    rows: list[dict[str, object]],
    *,
    query_target: str,
    headers: list[str] = HEADERS,
) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "检索结果"
    worksheet.append(headers)
    for row in rows:
        worksheet.append([row.get(header) for header in headers])
    workbook.create_sheet("获批适应症")
    conditions = workbook.create_sheet("检索条件")
    conditions.append(["NextPharma®"])
    conditions.append([f"检索结果:{len(rows)}条"])
    conditions.append(["创建时间:2026-07-16"])
    conditions.append([f"检索式: 靶点 = ({query_target})"])
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _version(
    session: Session,
    tenant: Tenant,
    tmp_path: Path,
    *,
    file_name: str,
    workbook: bytes,
    authorization_scopes: list[str],
) -> tuple[FileSystemObjectStore, SourceVersion]:
    store = FileSystemObjectStore(tmp_path / "objects")
    workbook_sha = hashlib.sha256(workbook).hexdigest()
    raw_object = store.put_bytes(tenant.id, "raw", workbook, workbook_sha, ".xlsx")
    extracted = b"[[sheet:test]]\n1\tplaceholder"
    extracted_sha = hashlib.sha256(extracted).hexdigest()
    extracted_object = store.put_bytes(tenant.id, "extracted-text", extracted, extracted_sha, ".txt")
    source = DataSource(
        tenant_id=tenant.id,
        name=f"Source {file_name}",
        source_type=DataSourceType.FOLDER,
        root_uri=str(tmp_path / file_name),
        owner="Pharma Data Operations",
        authorization_scopes=authorization_scopes,
        dataset_key="projects",
    )
    session.add(source)
    session.flush()
    asset = SourceAsset(
        tenant_id=tenant.id,
        data_source_id=source.id,
        logical_path=file_name,
        source_uri=str(tmp_path / file_name),
        file_name=file_name,
        extension=".xlsx",
        processing_mode="parse",
    )
    session.add(asset)
    session.flush()
    document = SourceDocument(
        tenant_id=tenant.id,
        title=file_name,
        source_type="folder",
        source_uri=asset.source_uri,
        content_sha256=workbook_sha,
    )
    session.add(document)
    session.flush()
    version = SourceVersion(
        tenant_id=tenant.id,
        source_asset_id=asset.id,
        version_number=1,
        content_sha256=workbook_sha,
        size_bytes=len(workbook),
        snapshot_status=StageStatus.SUCCEEDED,
        parse_status=StageStatus.SUCCEEDED,
        raw_object_uri=raw_object.uri,
        extracted_text_object_uri=extracted_object.uri,
        extracted_text_sha256=extracted_sha,
        source_document_id=document.id,
        state=SourceVersionState.PARSED,
    )
    session.add(version)
    session.commit()
    return store, version
